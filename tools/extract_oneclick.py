"""Read the supplied XLSM templates into offline browser data; never run VBA."""
import datetime as dt
import json
import pathlib
import re
import warnings

import openpyxl
from openpyxl.styles.colors import COLOR_INDEX as COLOR_INDEXED
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.utils.datetime import to_excel

ROOT = pathlib.Path(__file__).resolve().parents[1]
SOURCE = ROOT / "원클릭 프로그램(2026.5.수정)"
DEST = ROOT / "assets" / "oneclick-data.js"


def color(value, fallback):
    if value is None:
        return fallback
    if value.type == "rgb" and isinstance(value.rgb, str):
        return "#" + value.rgb[-6:]
    if value.type == "indexed" and value.indexed < len(COLOR_INDEXED):
        return "#" + COLOR_INDEXED[value.indexed][-6:]
    return fallback


def style(cell):
    font, border, align = cell.font, cell.border, cell.alignment
    css = {"fontSize": f"{font.sz or 11}pt", "fontWeight": "bold" if font.b else "normal",
           "color": color(font.color, "#111827"), "verticalAlign": align.vertical or "middle",
           "textAlign": {"centerContinuous": "center", "distributed": "left", "general": "left"}.get(align.horizontal, align.horizontal or "left")}
    if font.i:
        css["fontStyle"] = "italic"
    if font.u:
        css["textDecoration"] = "underline"
    if getattr(cell.fill, "patternType", None) == "solid":
        css["backgroundColor"] = color(cell.fill.fgColor, "#ffffff")
    for side in ("top", "right", "bottom", "left"):
        edge = getattr(border, side)
        if edge and edge.style:
            width = 2 if edge.style in ("medium", "thick", "double") else 1
            kind = "double" if edge.style == "double" else "dashed" if "dash" in edge.style.lower() else "solid"
            css["border" + side.title()] = f"{width}px {kind} {color(edge.color, '#222222')}"
    return {"css": css, "format": cell.number_format, "wrap": bool(align.wrap_text)}


def value(v):
    return to_excel(v) if isinstance(v, (dt.datetime, dt.date)) else v


def options(ws, address):
    for rule in ws.data_validations.dataValidation:
        if address in rule and rule.type == "list" and rule.formula1 and rule.formula1.startswith('"'):
            return [v.strip() for v in rule.formula1.strip('"').split(",")]
    return []


def input_fields(ws, kind):
    fields = []
    groups = (("발주기관", range(9, 16)), ("계약업체", range(19, 24)), ("계약·준공", range(27, 39))) if kind == "construction" else (("발주기관", range(10, 14)), ("계약업체", range(16, 20)), ("계약·완료", range(22, 28)))
    for group, rows in groups:
        for r in rows:
            for label_col, data_col in (("B", "C"), ("D", "E")):
                cell = ws[f"{data_col}{r}"]
                label = ws[f"{label_col}{r}"].value
                if not label or cell.data_type == "f" or isinstance(cell, openpyxl.cell.cell.MergedCell):
                    continue
                label = re.sub(r"\s+", "", str(label))
                typ = "date" if cell.is_date else "number" if isinstance(cell.value, (int, float)) and "번호" not in label else "text"
                if "율" in label and isinstance(cell.value, (int, float)):
                    typ = "percent"
                fields.append({"cell": cell.coordinate, "label": label, "group": group, "type": typ,
                               "options": options(ws, cell.coordinate), "example": value(cell.value)})
    return fields


def extract(path, key, kind, title):
    workbook = openpyxl.load_workbook(path, data_only=False)
    cached = openpyxl.load_workbook(path, data_only=True)
    styles, style_ids, sheets = [], {}, {}
    for ws in workbook:
        cells = {}
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None and not cell.has_style:
                    continue
                # References are retained, while large unused formatted regions are bounded.
                if cell.row > 300 or cell.column > 200:
                    continue
                skey = cell.style_id
                if skey not in style_ids:
                    style_ids[skey] = len(styles)
                    styles.append(style(cell))
                entry = {"s": style_ids[skey]}
                if cell.data_type == "f":
                    entry["f"] = cell.value
                    entry["cached"] = value(cached[ws.title][cell.coordinate].value)
                elif cell.value is not None:
                    entry["v"] = value(cell.value)
                cells[cell.coordinate] = entry
        printed = re.findall(r"\$?[A-Z]+\$?\d+:\$?[A-Z]+\$?\d+", str(ws.print_area or ""))
        bounds = printed[0].replace("$", "") if printed else f"A1:{get_column_letter(min(ws.max_column, 60))}{min(ws.max_row, 150)}"
        col_widths = {}
        for dimension in ws.column_dimensions.values():
            for column in range(dimension.min or 1, (dimension.max or dimension.min or 1) + 1):
                if column <= 200:
                    col_widths[str(column)] = round((dimension.width or 8.43) * 7 + 5, 2)
        sheets[ws.title] = {"cells": cells, "merges": [str(m) for m in ws.merged_cells.ranges],
                            "area": bounds, "widths": col_widths,
                            "heights": {str(r): d.height for r, d in ws.row_dimensions.items() if d.height},
                            "hiddenRows": [r for r, d in ws.row_dimensions.items() if d.hidden],
                            "state": ws.sheet_state, "landscape": ws.page_setup.orientation == "landscape"}
    forms = [ws.title for ws in workbook if re.match(r"^\d+\.", ws.title) and not ws.title.startswith(("1.", "2.")) and ws.sheet_state == "visible"] if kind != "fee" else ["설계용역비"]
    corrections = []
    if kind == "design":
        # The source affidavit accidentally links to the construction workbook.
        # Map by the printed labels (company, representative, agency, date).
        affidavit = sheets["9.수의계약 각서"]["cells"]
        replacements = {"H51": "='2.데이터입력'!C16", "J51": "='2.데이터입력'!E16&\"  (인)\"",
                        "C50": "='2.데이터입력'!C23",
                        "C52": affidavit["C52"]["f"].replace("'[2]2.데이터입력'!$C$10", "'2.데이터입력'!$C$11")}
        for address, formula in replacements.items():
            original = affidavit[address].get("f", affidavit[address].get("v"))
            corrections.append({"sheet": "9.수의계약 각서", "cell": address, "original": original, "formula": formula})
            affidavit[address] = {"s": affidavit[address]["s"], "f": formula}
    return {"key": key, "kind": kind, "title": title, "file": str(path.relative_to(ROOT)).replace("\\", "/"),
            "styles": styles, "sheets": sheets, "forms": forms,
            "corrections": corrections,
            "fields": input_fields(workbook["2.데이터입력"], kind) if kind != "fee" else []}


def main():
    specs = [("construction-seoul", "construction", "공사서류 · 서울시교육청"),
             ("construction-other", "construction", "공사서류 · 타기관"),
             ("design-seoul", "design", "설계용역서류 · 서울시교육청"),
             ("design-other", "design", "설계용역서류 · 타기관"),
             ("design-fee", "fee", "설계용역비 산출")]
    result = {}
    for path, (key, kind, title) in zip(sorted(SOURCE.glob("*.xlsm")), specs, strict=True):
        result[key] = extract(path, key, kind, title)
        print(key, "forms", len(result[key]["forms"]), "fields", len(result[key]["fields"]))
    DEST.parent.mkdir(exist_ok=True)
    DEST.write_text("/* Generated from the supplied 2026.5 XLSM files. Do not edit by hand. */\nwindow.ONECLICK_BOOKS = " + json.dumps(result, ensure_ascii=False, separators=(",", ":")) + ";\n", encoding="utf-8")
    print("Wrote", DEST, DEST.stat().st_size, "bytes")


if __name__ == "__main__":
    warnings.simplefilter("ignore", UserWarning)
    main()
